import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"D:\Prototype CAS\EDIT MASTER MASAL")
OUT_DIR = ROOT / "outputs" / "019fe01d-d395-7bd1-91da-86a66500c3bd"
FINAL = OUT_DIR / "template-import-pengguna_batch2_siap_upload.xlsx"
TEMPLATE = Path(r"C:\Users\KOKO98\Downloads\template-import-pengguna (1).xlsx")
PREPARED = OUT_DIR / "batch2_user_analysis" / "prepared_user_import.json"
REPORT = OUT_DIR / "batch2_user_verification.json"

prepared = json.loads(PREPARED.read_text(encoding="utf-8"))
expected_users = prepared["users"]
expected_satkers = prepared["satkers"]

wb = load_workbook(FINAL, read_only=True, data_only=False)
assert wb.sheetnames == ["Data Pengguna", "Daftar Satker", "Petunjuk"], wb.sheetnames

data = wb["Data Pengguna"]
satker = wb["Daftar Satker"]
petunjuk = wb["Petunjuk"]
expected_headers = ["Nama", "NRP/NIP", "Email", "Email Dinas Polri", "Username", "Role", "Kode atau Nama Satker", "Status"]
actual_headers = [data.cell(1, col).value for col in range(1, 9)]
assert actual_headers == expected_headers, actual_headers

actual_users = []
for row in data.iter_rows(min_row=2, max_col=8, values_only=True):
    if not any(value not in (None, "") for value in row):
        continue
    actual_users.append(["" if value is None else str(value) for value in row])
expected_user_rows = [[
    item["Nama"], ("'" + item["NRP/NIP"] if item["NRP/NIP"].startswith("0") else item["NRP/NIP"]), item["Email"], item["Email Dinas Polri"],
    item["Username"], item["Role"], item["Kode atau Nama Satker"], item["Status"],
] for item in expected_users]
assert actual_users == expected_user_rows
assert len(actual_users) == 104

names = [row[0] for row in actual_users]
nrps = [row[1] for row in actual_users]
emails = [row[2].lower() for row in actual_users]
email_dinas = [row[3].lower() for row in actual_users]
roles = [row[5] for row in actual_users]
user_codes = [row[6] for row in actual_users]
statuses = [row[7] for row in actual_users]
assert all(names) and all(nrps) and all(emails) and all(email_dinas) and all(user_codes)
assert all(row[4] == "" for row in actual_users)
assert sum(nrp.startswith("'0") for nrp in nrps) == 24
assert len(set(nrps)) == len(nrps)
assert len(set(emails)) == len(emails)
assert len(set(email_dinas)) == len(email_dinas)
assert Counter(roles) == {"operator": 104}
assert Counter(statuses) == {"AKTIF": 104}

actual_satkers = []
for row in satker.iter_rows(min_row=2, max_col=3, values_only=True):
    if not any(value not in (None, "") for value in row):
        continue
    actual_satkers.append(["" if value is None else str(value) for value in row])
expected_satker_rows = [[
    item["Kode"], item["Nama Satker"], item["Jalur (Polda > Polres > Polsek > Satuan)"],
] for item in expected_satkers]
assert actual_satkers == expected_satker_rows
assert len(actual_satkers) == 13865
satker_codes = [row[0] for row in actual_satkers]
assert len(set(satker_codes)) == len(satker_codes)
assert set(user_codes).issubset(set(satker_codes))
assert all(code.startswith("060") for code in satker_codes)

changed_users = {
    item["Nama"]: (item["kode_asli"], item["Kode atau Nama Satker"])
    for item in expected_users if item["kode_asli"] != item["Kode atau Nama Satker"]
}
assert len(changed_users) == 6
actual_by_name = {row[0]: row for row in actual_users}
for name, (old_code, new_code) in changed_users.items():
    assert actual_by_name[name][6] == new_code
    assert actual_by_name[name][6] != old_code

pbd_expected = {
    "POLRESTA SORONG KOTA": "060.01.36.101",
    "POLRES SORONG": "060.01.36.102",
    "POLRES SORONG SELATAN": "060.01.36.103",
    "POLRES RAJA AMPAT": "060.01.36.104",
    "POLRES MAYBRAT": "060.01.36.105",
    "POLRES TAMBRAUW": "060.01.36.106",
}
pbd_rows = {}
for code, name, path in actual_satkers:
    if name in pbd_expected:
        pbd_rows[name] = {"code": code, "path": path}
for name, code in pbd_expected.items():
    assert pbd_rows[name]["code"] == code, (name, pbd_rows[name])
    assert pbd_rows[name]["path"].startswith("POLDA PAPUA BARAT DAYA"), (name, pbd_rows[name])

template_wb = load_workbook(TEMPLATE, read_only=True, data_only=False)
template_petunjuk = list(template_wb["Petunjuk"].iter_rows(min_row=1, max_row=12, max_col=2, values_only=True))
final_petunjuk = list(petunjuk.iter_rows(min_row=1, max_row=12, max_col=2, values_only=True))
normalize_blank = lambda rows: [["" if value is None else value for value in row] for row in rows]
assert normalize_blank(final_petunjuk) == normalize_blank(template_petunjuk)

formula_errors = []
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}:
                formula_errors.append(f"{sheet.title}!{cell.coordinate}:{cell.value}")
assert not formula_errors

report = {
    "file": str(FINAL),
    "sheet_names": wb.sheetnames,
    "users": len(actual_users),
    "unique_nrp_nip": len(set(nrps)),
    "unique_email": len(set(emails)),
    "unique_email_dinas": len(set(email_dinas)),
    "role_counts": dict(Counter(roles)),
    "status_counts": dict(Counter(statuses)),
    "blank_usernames": sum(row[4] == "" for row in actual_users),
    "nrp_with_leading_zero_apostrophe": sum(nrp.startswith("'0") for nrp in nrps),
    "all_user_satker_codes_exist": True,
    "changed_user_codes": changed_users,
    "satker_directory_rows": len(actual_satkers),
    "unique_satker_codes": len(set(satker_codes)),
    "papua_barat_daya": pbd_rows,
    "petunjuk_unchanged": True,
    "formula_errors": formula_errors,
    "excluded_duplicate_source_rows": prepared["summary"]["excluded_duplicate_rows"],
    "source_notes_reviewed": 12,
}
REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False, indent=2))
