import json
import sys
from pathlib import Path

from docx import Document


def iter_block_items(parent):
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P

    parent_elm = parent.element.body
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def main():
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])
    doc = Document(src)
    blocks = []
    for block in iter_block_items(doc):
        if hasattr(block, "rows"):
            rows = []
            for row in block.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            blocks.append({"type": "table", "rows": rows})
        else:
            text = block.text.strip()
            if text:
                blocks.append(
                    {
                        "type": "paragraph",
                        "style": block.style.name if block.style else "",
                        "text": text,
                    }
                )
    payload = {
        "source": str(src),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "section_count": len(doc.sections),
        "blocks": blocks,
    }
    dst.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
