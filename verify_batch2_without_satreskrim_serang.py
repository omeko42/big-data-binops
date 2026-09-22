import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"D:\Prototype CAS\EDIT MASTER MASAL")
OUT_DIR = ROOT / "outputs" / "019fe01d-d395-7bd1-91da-86a66500c3bd"
SOURCE = OUT_DIR / "import-pengguna_batch2_siap_upload.xlsx"
FINAL = OUT_DIR / "import-pengguna_batch2_siap_upload_tanpa_satreskrim_serang.xlsx"
REPORT = OUT_DIR / "batch2_tanpa_satreskrim_serang_verification.json"
EXCLUDED_CODE = "060.01.14.02.00.10"

wb = load_workbook(FINAL, read_only=False, data_only=False)
source_wb = load_workbook(SOURCE, read_only=True, data_only=False)
assert wb.sheetnames == ["Data Pengguna", "Daftar Satker", "Petunjuk"]
data = wb["Data Pengguna"]
satker = wb["Daftar Satker"]

headers = [data.cell(1, col).value for col in range(1, 9)]
assert headers == ["Nama", "NRP/NIP", "Email", "Email Dinas Polri", "Username", "Role", "Kode atau Nama Satker", "Status"]
rows = []
for row in data.iter_rows(min_row=2, max_col=8, values_only=True):
    values = ["" if value is None else str(value) for value in row]
    if any(values):
        rows.append(values)
assert len(rows) == 59
assert all(row[6] != EXCLUDED_CODE for row in rows)
assert all(row[4] == "" for row in rows)
assert Counter(row[5] for row in rows) == {"operator": 59}
assert Counter(row[7] for row in rows) == {"AKTIF": 59}
assert len({row[1] for row in rows}) == 59
assert len({row[2].lower() for row in rows}) == 59
assert len({row[3].lower() for row in rows}) == 59

satker_codes = {
    str(row[0]) for row in satker.iter_rows(min_row=2, max_col=1, values_only=True) if row[0] not in (None, "")
}
assert len(satker_codes) == 13865
assert {row[6] for row in rows}.issubset(satker_codes)

assert len(data.tables) == 1
table = list(data.tables.values())[0]
assert table.ref == "A1:H60", table.ref

def normalize(rows):
    return [["" if value is None else value for value in row] for row in rows]

for sheet_name, max_col in [("Daftar Satker", 3), ("Petunjuk", 2)]:
    final_rows = list(wb[sheet_name].iter_rows(max_col=max_col, values_only=True))
    source_rows = list(source_wb[sheet_name].iter_rows(max_col=max_col, values_only=True))
    assert normalize(final_rows) == normalize(source_rows)

errors = []
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}:
                errors.append(f"{sheet.title}!{cell.coordinate}:{cell.value}")
assert not errors

source_data = source_wb["Data Pengguna"]
source_count = sum(1 for row in source_data.iter_rows(min_row=2, max_col=8, values_only=True) if any(value not in (None, "") for value in row))
source_serang_count = sum(
    1 for row in source_data.iter_rows(min_row=2, max_col=8, values_only=True)
    if row[6] is not None and str(row[6]) == EXCLUDED_CODE
)
assert source_count == 104
assert source_serang_count == 45

report = {
    "file": str(FINAL),
    "source_users": source_count,
    "excluded_satreskrim_serang": source_serang_count,
    "remaining_users": len(rows),
    "excluded_code_absent": True,
    "roles": dict(Counter(row[5] for row in rows)),
    "statuses": dict(Counter(row[7] for row in rows)),
    "blank_usernames": sum(row[4] == "" for row in rows),
    "unique_nrp_nip": len({row[1] for row in rows}),
    "unique_email": len({row[2].lower() for row in rows}),
    "all_user_codes_exist": True,
    "table_range": table.ref,
    "daftar_satker_unchanged": True,
    "petunjuk_unchanged": True,
    "formula_errors": errors,
}
REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False, indent=2))
