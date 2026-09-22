import csv
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"D:\Prototype CAS\EDIT MASTER MASAL")
BATCH = Path(r"C:\Users\KOKO98\Downloads\batch 2 pengguna cas.xlsx")
TEMPLATE = Path(r"C:\Users\KOKO98\Downloads\template-import-pengguna (1).xlsx")
RAW = ROOT / "bdb_satker_raw.csv"
V4_MAP = ROOT / "bdb_satker_v4_map.json"
OUT_DIR = ROOT / "outputs" / "019fe01d-d395-7bd1-91da-86a66500c3bd" / "batch2_user_analysis"
OUT_DIR.mkdir(parents=True, exist_ok=True)


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def norm(value):
    value = clean(value).upper()
    value = re.sub(r"[^A-Z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def nrp_text(value):
    text = clean(value)
    return re.sub(r"\.0$", "", text)


wb = load_workbook(BATCH, read_only=True, data_only=True)
resp = wb["RESPONS CAS"]
prepared = wb["SIAP IMPOR CAS"]
resp_rows = list(resp.iter_rows(values_only=True))
prepared_rows = list(prepared.iter_rows(values_only=True))

template_wb = load_workbook(TEMPLATE, read_only=True, data_only=True)
satker_sheet = template_wb["Daftar Satker"]
template_satker = {}
name_to_template = defaultdict(list)
path_to_template = defaultdict(list)
for row in satker_sheet.iter_rows(min_row=2, values_only=True):
    code, name, path = (clean(row[0]), clean(row[1]), clean(row[2]))
    if not code:
        continue
    template_satker[code] = {"code": code, "name": name, "path": path}
    name_to_template[norm(name)].append(template_satker[code])
    path_to_template[norm(path)].append(template_satker[code])

with RAW.open("r", encoding="utf-8-sig", newline="") as handle:
    raw_rows = list(csv.DictReader(handle))
with V4_MAP.open("r", encoding="utf-8") as handle:
    v4_map = json.load(handle)

by_raw_code = defaultdict(list)
by_raw_name = defaultdict(list)
by_id = {}
for item in raw_rows:
    by_id[item["id"]] = item
    by_raw_code[clean(item["kode"])].append(item)
    by_raw_code[clean(item["kode_canonical"])].append(item)
    by_raw_name[norm(item["deskripsi"])].append(item)


def map_code(original_code, chosen_name, chosen_path):
    candidates = by_raw_code.get(original_code, [])
    if len(candidates) > 1:
        exact_name = [item for item in candidates if norm(item["deskripsi"]) == norm(chosen_name)]
        if exact_name:
            candidates = exact_name
    if not candidates:
        named = by_raw_name.get(norm(chosen_name), [])
        if len(named) == 1:
            candidates = named
        elif norm(chosen_path) in path_to_template:
            path_codes = {item["code"] for item in path_to_template[norm(chosen_path)]}
            candidates = [item for code in path_codes for item in by_raw_code.get(code, [])]
    if not candidates:
        return None, "TIDAK_DITEMUKAN", []
    item = candidates[0]
    mapped = clean(v4_map.get(item["id"], {}).get("kode_canonical_v4")) or clean(item["kode_canonical"]) or clean(item["kode"])
    status = "KODE_DIUBAH_V4" if mapped != original_code else "KODE_TETAP"
    return mapped, status, candidates


audits = []
for excel_row, (source, ready) in enumerate(zip(resp_rows[1:], prepared_rows[1:]), start=2):
    nama = clean(source[1] or ready[0])
    nrp = nrp_text(source[2] or ready[1])
    email = clean(source[3] or ready[2]).lower()
    email_dinas = clean(ready[3] or source[3]).lower()
    username = clean(ready[4])
    role = clean(ready[5] or "operator").lower()
    original_code = clean(source[6] or ready[6])
    chosen_name = clean(source[7] or ready[12])
    chosen_path = clean(source[5] or ready[13])
    keterangan = clean(source[8])
    daftar = clean(source[10] or ready[8])
    status = clean(ready[7] or "AKTIF").upper()
    mapped_code, map_status, candidates = map_code(original_code, chosen_name, chosen_path)
    template_entry = template_satker.get(original_code)
    selected_name_match = bool(template_entry and norm(template_entry["name"]) == norm(chosen_name))
    selected_path_match = bool(template_entry and (
        norm(template_entry["path"]) == norm(chosen_path)
        or norm(chosen_path).replace(" > ", " ") == norm(template_entry["path"])
    ))
    flags = []
    if keterangan:
        flags.append("KETERANGAN_PERLU_DIBACA")
    if not mapped_code:
        flags.append("SATKER_TIDAK_DITEMUKAN")
    if template_entry and not selected_name_match:
        flags.append("NAMA_TIDAK_COCOK_DENGAN_KODE")
    if template_entry and not selected_path_match:
        flags.append("JALUR_TIDAK_PERSIS_DENGAN_TEMPLATE")
    if not nama:
        flags.append("NAMA_KOSONG")
    if not email or "@" not in email:
        flags.append("EMAIL_TIDAK_VALID")
    if role not in {"superadmin", "admin_binops", "operator", "viewer"}:
        flags.append("ROLE_TIDAK_VALID")
    if status not in {"AKTIF", "NONAKTIF"}:
        flags.append("STATUS_TIDAK_VALID")
    audits.append({
        "excel_row": excel_row,
        "nama": nama,
        "nrp_nip": nrp,
        "email": email,
        "email_dinas": email_dinas,
        "username": username,
        "role": role,
        "kode_asli": original_code,
        "kode_final": mapped_code or "",
        "status_kode": map_status,
        "nama_satker_respons": chosen_name,
        "jalur_satker_respons": chosen_path,
        "keterangan_respons": keterangan,
        "daftar_respons": daftar,
        "status": status,
        "template_name": template_entry["name"] if template_entry else "",
        "template_path": template_entry["path"] if template_entry else "",
        "flags": flags,
        "candidate_count": len(candidates),
    })

for field, label in [("email", "DUPLIKAT_EMAIL_BATCH"), ("email_dinas", "DUPLIKAT_EMAIL_DINAS_BATCH"), ("nrp_nip", "DUPLIKAT_NRP_BATCH")]:
    groups = defaultdict(list)
    for item in audits:
        if item[field]:
            groups[item[field]].append(item)
    for group in groups.values():
        if len(group) > 1:
            for item in group:
                item["flags"].append(label)

summary = {
    "batch_rows": len(audits),
    "ready_rows": sum(not item["flags"] for item in audits),
    "rows_with_flags": sum(bool(item["flags"]) for item in audits),
    "rows_with_response_note": sum(bool(item["keterangan_respons"]) for item in audits),
    "codes_changed_by_v4": sum(item["status_kode"] == "KODE_DIUBAH_V4" for item in audits),
    "codes_not_found": sum(not item["kode_final"] for item in audits),
    "flag_counts": dict(Counter(flag for item in audits for flag in item["flags"])),
    "response_notes": [
        {key: item[key] for key in ["excel_row", "nama", "nrp_nip", "kode_asli", "kode_final", "nama_satker_respons", "jalur_satker_respons", "keterangan_respons", "flags"]}
        for item in audits if item["keterangan_respons"]
    ],
    "changed_codes": [
        {key: item[key] for key in ["excel_row", "nama", "kode_asli", "kode_final", "nama_satker_respons", "jalur_satker_respons", "keterangan_respons"]}
        for item in audits if item["status_kode"] == "KODE_DIUBAH_V4"
    ],
}


def identity_score(item):
    local = item["email"].split("@", 1)[0]
    score = 0
    if item["email"].endswith("@polri.go.id"):
        score += 1
    if local.isdigit():
        if local == item["nrp_nip"]:
            score += 5
        elif local.lstrip("0") == item["nrp_nip"].lstrip("0"):
            score += 4
    return score


# Buang hanya kiriman ganda dari orang yang sama. NRP yang sama tetapi nama berbeda
# tidak otomatis dibuang karena dapat berasal dari angka Excel yang rusak.
identity_groups = defaultdict(list)
for item in audits:
    identity_groups[(norm(item["nama"]), item["nrp_nip"])].append(item)

kept = []
excluded_duplicates = []
for group in identity_groups.values():
    ordered = sorted(group, key=lambda item: (-identity_score(item), item["excel_row"]))
    kept.append(ordered[0])
    excluded_duplicates.extend(ordered[1:])
kept.sort(key=lambda item: item["excel_row"])

prepared_users = []
for item in kept:
    local = item["email"].split("@", 1)[0]
    if local.isdigit():
        final_nrp = local
        nrp_source = "LOCAL_PART_EMAIL_DINAS"
    else:
        final_nrp = item["nrp_nip"].zfill(8) if item["nrp_nip"].isdigit() and len(item["nrp_nip"]) < 8 else item["nrp_nip"]
        nrp_source = "NRP_BATCH_DENGAN_PADDING" if final_nrp != item["nrp_nip"] else "NRP_BATCH"
    prepared_users.append({
        "Nama": item["nama"],
        "NRP/NIP": final_nrp,
        "Email": item["email"],
        "Email Dinas Polri": item["email_dinas"],
        "Username": item["username"],
        "Role": item["role"],
        "Kode atau Nama Satker": item["kode_final"],
        "Status": item["status"],
        "source_row": item["excel_row"],
        "kode_asli": item["kode_asli"],
        "nrp_asli_excel": item["nrp_nip"],
        "nrp_source": nrp_source,
        "keterangan_respons": item["keterangan_respons"],
    })

updated_satkers = []
unmapped_satkers = []
satker_paths_changed = 0
for code, entry in template_satker.items():
    candidates = by_raw_code.get(code, [])
    exact = [item for item in candidates if norm(item["deskripsi"]) == norm(entry["name"])]
    if exact:
        candidates = exact
    if candidates:
        item = candidates[0]
        final_code = clean(v4_map.get(item["id"], {}).get("kode_canonical_v4")) or clean(item["kode_canonical"]) or code
    else:
        final_code = code
        unmapped_satkers.append(entry)
    final_path = entry["path"]
    # Unit 101-106 sudah ditempatkan di bawah POLDA PAPUA BARAT DAYA pada kode V4.
    # Selaraskan label jalurnya juga agar pencarian nama/jalur tidak mengarah ke polda lama.
    if re.match(r"^060\.01\.36\.(101|102|103|104|105|106)(?:\.|$)", final_code):
        old_prefix = "POLDA PAPUA BARAT >"
        if final_path.startswith(old_prefix):
            final_path = "POLDA PAPUA BARAT DAYA >" + final_path[len(old_prefix):]
    if final_path != entry["path"]:
        satker_paths_changed += 1
    updated_satkers.append({
        "Kode": final_code,
        "Nama Satker": entry["name"],
        "Jalur (Polda > Polres > Polsek > Satuan)": final_path,
        "kode_asli_template": code,
    })

prepared_summary = {
    "source_rows": len(audits),
    "prepared_users": len(prepared_users),
    "excluded_duplicate_rows": [item["excel_row"] for item in excluded_duplicates],
    "nrp_restored_from_email": sum(item["nrp_source"] == "LOCAL_PART_EMAIL_DINAS" and item["NRP/NIP"] != item["nrp_asli_excel"] for item in prepared_users),
    "nrp_padded_from_batch": sum(item["nrp_source"] == "NRP_BATCH_DENGAN_PADDING" for item in prepared_users),
    "user_codes_changed": sum(item["kode_asli"] != item["Kode atau Nama Satker"] for item in prepared_users),
    "satker_directory_rows": len(updated_satkers),
    "satker_directory_codes_changed": sum(item["Kode"] != item["kode_asli_template"] for item in updated_satkers),
    "satker_directory_paths_changed": satker_paths_changed,
    "satker_directory_unmapped": len(unmapped_satkers),
    "unique_emails": len({item["Email"] for item in prepared_users}),
    "unique_email_dinas": len({item["Email Dinas Polri"] for item in prepared_users}),
    "unique_final_nrp": len({item["NRP/NIP"] for item in prepared_users}),
    "unique_final_codes": len({item["Kode"] for item in updated_satkers}),
}

with (OUT_DIR / "analysis.json").open("w", encoding="utf-8") as handle:
    json.dump({"summary": summary, "prepared_summary": prepared_summary, "rows": audits}, handle, ensure_ascii=False, indent=2)
with (OUT_DIR / "prepared_user_import.json").open("w", encoding="utf-8") as handle:
    json.dump({
        "summary": prepared_summary,
        "users": prepared_users,
        "satkers": updated_satkers,
        "excluded_duplicates": [
            {key: item[key] for key in ["excel_row", "nama", "nrp_nip", "email", "kode_final", "nama_satker_respons"]}
            for item in excluded_duplicates
        ],
        "unmapped_satkers": unmapped_satkers,
    }, handle, ensure_ascii=False, indent=2)
with (OUT_DIR / "audit.csv").open("w", encoding="utf-8-sig", newline="") as handle:
    writer = csv.DictWriter(handle, fieldnames=[
        "excel_row", "nama", "nrp_nip", "email", "email_dinas", "username", "role",
        "kode_asli", "kode_final", "status_kode", "nama_satker_respons", "jalur_satker_respons",
        "keterangan_respons", "daftar_respons", "status", "template_name", "template_path", "flags",
    ])
    writer.writeheader()
    for item in audits:
        row = {key: item.get(key, "") for key in writer.fieldnames}
        row["flags"] = " | ".join(item["flags"])
        writer.writerow(row)

print(json.dumps({"analysis": summary, "prepared": prepared_summary}, ensure_ascii=False, indent=2))
